#!/usr/bin/env python3
"""Build the drafts site.

Reads src/content/*.md, src/template.html and src/VERSION; writes
<site>/index.html (current), <site>/versions/v<ver>/index.html (frozen copy),
and <site>/downloads/<base>.pdf, <base>-with-sources.pdf, <base>.docx, <base>.md
for each deliverable (copied under the frozen version too).

Markdown subset: #/##/### headings, paragraphs, - bullets, 1. numbered,
| tables |, > blockquotes, **bold**, *italic*, `code`, ---.
Provenance: lines starting with "::src " directly under a heading:
    ::src Label :: explanation
Blocks:
    :::details Title ... :::      collapsible block (headings inside stay inside)
    :::register ... :::           a decisions table with a person filter; the row's
                                  first cell is its id; the Owner column drives the filter
    :::why ID ... :::             details for register row ID (shown on hover/click)
"""
import html, re, pathlib, shutil, subprocess, tempfile, sys

ROOT = pathlib.Path(__file__).parent          # src/
SITE = ROOT.parent                             # site root

DOCS = [
    # key, tab title, content file, group, download basename (None = no downloads)
    ("case",      "Case-making",          "01-case-making.md",          "deliverable", "chss-case-making"),
    ("district",  "District commitments", "02-district-commitments.md", "deliverable", "chss-district-commitments"),
    ("chss",      "CHSS commitment",      "03-chss-commitment.md",      "deliverable", "chss-privacy-security-commitment"),
    ("notes",     "Notes",                "04-notes-for-james.md",      "note",        None),
    ("decisions", "Decisions",            "05-decisions.md",            "note",        None),
    ("changelog", "Changelog",            "06-changelog.md",            "note",        None),
]

LABEL_CLASS = {
    "Luke's experience": "src-luke", "AFUHSD precedent": "src-luke", "Luke's decision": "src-luke",
    "Conversation": "src-call",
    "James's collateral": "src-james", "James's MOU draft": "src-james",
    "Emily's architecture": "src-emily", "Emily's email": "src-emily",
    "FERPA": "src-law", "PPRA": "src-law", "COPPA": "src-law", "State law": "src-law", "Case law / guidance": "src-law",
    "SDPC NDPA": "src-sdpc",
    "Columbus City DSA": "src-district", "San Diego CCAP": "src-district", "Lewis County": "src-district", "District lens": "src-district",
    "af-dpa": "src-afdpa",
    "Drafting choice": "src-choice", "Needs CHSS confirmation": "src-other",
}
PEOPLE = ["Jonah", "Emily", "Tara", "Demian", "James", "Counsel"]

HOWTO = (
    '<div class="srcwrap howto"><button class="pillbtn" type="button" aria-expanded="false" '
    'aria-label="How to read this page">How to read this</button>'
    '<div class="pop" role="dialog"><div class="pop-h">How to read this</div>'
    '<p>Working drafts for the accelerator agreement. Every section has a <span class="tag src-choice">Source</span> '
    'pill in the right-hand column; hover or tap it to see where the language came from.</p>'
    '<p><span class="tag src-luke">Luke\'s experience</span><span class="tag src-call">Conversation</span>'
    '<span class="tag src-james">James\'s collateral</span><span class="tag src-emily">Emily\'s architecture</span>'
    '<span class="tag src-law">FERPA / PPRA / state law</span><span class="tag src-sdpc">SDPC NDPA</span>'
    '<span class="tag src-district">District agreements / district lens</span><span class="tag src-afdpa">af-dpa knowledge base</span>'
    '<span class="tag src-luke">Luke\'s decision</span></p>'
    '<p>For discussion among the CHSS team and partner districts. Not legal advice; counsel reviews before anything is signed.</p>'
    '</div></div>'
)

BLOCK_START = re.compile(r"^(#{1,3} |::src |- |\d+\. |\| |> |---$|:::)")


# ---------------------------------------------------------------- inline / sources
def inline(s):
    s = html.escape(s, quote=False)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", s)
    s = re.sub(r"`(.+?)`", r"<code>\1</code>", s)
    return s


def sources_list(srcs):
    return "".join(
        f'<li><span class="tag {LABEL_CLASS.get(label, "src-other")}">{html.escape(label)}</span> {inline(text)}</li>'
        for label, text in srcs)


def render_sources(srcs):
    if not srcs:
        return ""
    return ('<div class="srcwrap"><button class="pillbtn" type="button" aria-expanded="false" '
            'aria-label="Where this language came from">Source</button>'
            f'<div class="pop" role="dialog"><div class="pop-h">Where this came from</div><ul>{sources_list(srcs)}</ul></div></div>')


def parse_table(table_lines):
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table_lines]
    return [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]


def table_html(cells):
    if not cells:
        return "<table></table>"
    t = ["<table>", "<thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in cells[0]) + "</tr></thead><tbody>"]
    for r in cells[1:]:
        t.append("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>")
    t.append("</tbody></table>")
    return "".join(t)


# ---------------------------------------------------------------- register
def who_tokens(owner_text):
    low = owner_text.lower()
    return [p.lower() for p in PEOPLE if p.lower() in low]


def render_register(table_lines, whys, reg_id):
    rows = parse_table(table_lines)
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    owner_ix = next((ix for ix, h in enumerate(head) if h.strip().lower() == "owner"), None)
    people_seen, out = [], []
    for r in body:
        rid = r[0].strip()
        who = who_tokens(r[owner_ix]) if owner_ix is not None and owner_ix < len(r) else []
        for w in who:
            if w not in people_seen:
                people_seen.append(w)
        cells = []
        for ci, c in enumerate(r):
            if ci == 1 and rid in whys:
                cells.append(f'<td><span class="regwrap">{inline(c)} '
                             f'<button class="pillbtn why" type="button" aria-expanded="false" aria-label="Why">Why</button>'
                             f'<div class="pop" role="dialog">{whys[rid]}</div></span></td>')
            else:
                cells.append(f"<td>{inline(c)}</td>")
        out.append(f'<tr data-who="{" ".join(who)}">' + "".join(cells) + "</tr>")
    chips = ['<button class="chip" type="button" data-who="all" aria-pressed="true">Everyone</button>']
    chips += [f'<button class="chip" type="button" data-who="{w}" aria-pressed="false">{w.capitalize()}</button>' for w in people_seen]
    thead = "<thead><tr>" + "".join(f"<th>{inline(h)}</th>" for h in head) + "</tr></thead>"
    return (f'<div class="regfilter" data-register="{reg_id}"><span class="regfilter-l">Show for</span>{"".join(chips)}</div>'
            f'<div class="tablewrap"><table class="register" id="{reg_id}">{thead}<tbody>{"".join(out)}</tbody></table></div>')


def render_why(md):
    srcs, body = [], []
    for ln in md.splitlines():
        m = re.match(r"^::src (.+?) :: (.*)$", ln)
        if m:
            srcs.append((m.group(1).strip(), m.group(2).strip()))
        else:
            body.append(ln)
    h = '<div class="pop-h">Why, and what settles it</div>' + convert("\n".join(body), plain=True)
    if srcs:
        h += f'<div class="pop-h" style="margin-top:8px">Sources</div><ul>{sources_list(srcs)}</ul>'
    return h


# ---------------------------------------------------------------- converter
def convert(md, plain=False, title_extra=""):
    whys, lines, src_lines = {}, [], md.splitlines()
    i = 0
    while i < len(src_lines):
        m = re.match(r"^:::why\s+(\S+)\s*$", src_lines[i])
        if m:
            buf = []
            i += 1
            while i < len(src_lines) and src_lines[i].strip() != ":::":
                buf.append(src_lines[i]); i += 1
            whys[m.group(1)] = render_why("\n".join(buf))
            i += 1
            continue
        lines.append(src_lines[i]); i += 1

    out, list_stack = [], []
    i = 0
    pending = None
    sec_open = False
    in_fold = 0
    doc_srcs = None
    first_block = False
    reg_count = 0

    def close_lists():
        while list_stack:
            out.append(f"</{list_stack.pop()}>")

    def emit_block(h):
        nonlocal first_block
        if first_block:
            out.append(f'<div class="firstwrap">{h}{render_sources(doc_srcs)}</div>')
            first_block = False
        else:
            out.append(h)

    def flush_heading():
        nonlocal pending, sec_open, doc_srcs, first_block
        if pending is None:
            return
        level, text, srcs = pending
        close_lists()
        if plain or in_fold:
            lvl = min(level + (2 if in_fold else 1), 6)
            out.append(f'<div class="sec-head inner"><h{lvl}>{inline(text)}</h{lvl}>{render_sources(srcs)}</div>')
        elif level >= 2:
            if first_block:
                out.append(render_sources(doc_srcs)); first_block = False
            if sec_open:
                out.append("</section>")
            sid = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
            out.append(f'<section class="sec lvl{level}" id="{sid}">')
            sec_open = True
            out.append(f'<div class="sec-head"><h{level}>{inline(text)}</h{level}>{render_sources(srcs)}</div>')
        else:
            out.append(f'<div class="titlewrap"><h1>{inline(text)}</h1>{HOWTO}</div>')
            if title_extra:
                out.append(title_extra)
            doc_srcs = srcs
            first_block = bool(srcs)
        pending = None

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("::src "):
            if pending is not None:
                label, _, text = ln[6:].partition(" :: ")
                pending[2].append((label.strip(), text.strip()))
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            flush_heading()
            pending = (len(m.group(1)), m.group(2).strip(), [])
            i += 1
            continue
        flush_heading()
        if not ln.strip():
            close_lists(); i += 1; continue
        if ln.strip() == "---":
            close_lists(); out.append("<hr>"); i += 1; continue
        md_ = re.match(r"^:::details\s+(.*)$", ln)
        if md_:
            close_lists()
            out.append(f'<details class="fold"><summary>{inline(md_.group(1).strip())}</summary><div class="fold-body">')
            in_fold += 1
            i += 1
            continue
        if ln.strip() == ":::register":
            close_lists()
            i += 1
            tbl = []
            while i < len(lines) and lines[i].strip() != ":::":
                if lines[i].startswith("|"):
                    tbl.append(lines[i])
                i += 1
            i += 1
            reg_count += 1
            emit_block(render_register(tbl, whys, f"register-{reg_count}"))
            continue
        if ln.strip() == ":::":
            close_lists()
            if in_fold:
                out.append("</div></details>"); in_fold -= 1
            i += 1
            continue
        if ln.startswith("|"):
            close_lists()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            emit_block(table_html(parse_table(rows)))
            continue
        if ln.startswith("> "):
            close_lists()
            q = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i][1:].strip()); i += 1
            paras = "\n".join(q).split("\n\n")
            emit_block("<blockquote>" + "".join(f"<p>{inline(p.replace(chr(10), ' '))}</p>" for p in paras if p.strip()) + "</blockquote>")
            continue
        mb = re.match(r"^(\s*)- (.*)$", ln)
        mo = re.match(r"^(\s*)\d+\. (.*)$", ln)
        if mb or mo:
            kind = "ul" if mb else "ol"
            if not list_stack or list_stack[-1] != kind:
                close_lists(); out.append(f"<{kind}>"); list_stack.append(kind)
            out.append(f"<li>{inline((mb or mo).group(2))}</li>")
            i += 1
            continue
        p = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            p.append(lines[i].strip()); i += 1
        close_lists()
        emit_block(f"<p>{inline(' '.join(p))}</p>")
    flush_heading()
    close_lists()
    while in_fold:
        out.append("</div></details>"); in_fold -= 1
    if first_block:
        out.append(render_sources(doc_srcs))
    if sec_open:
        out.append("</section>")
    return "\n".join(out)


# ---------------------------------------------------------------- print / downloads
PRINT_CSS = """
body{font-family:"Iowan Old Style","Palatino Linotype",Palatino,Georgia,serif;font-size:11pt;line-height:1.5;color:#111;margin:0;padding:0 .2in}
h1{font-size:20pt;margin:0 0 .2em}h2{font-size:14.5pt;margin:1.4em 0 .3em}h3{font-size:12pt;margin:1.1em 0 .3em}
p{margin:.5em 0}ul,ol{margin:.4em 0 .7em;padding-left:1.3em}li{margin:.2em 0}
table{border-collapse:collapse;width:100%;margin:.6em 0 1em;font-size:9.5pt;font-family:-apple-system,Helvetica,Arial,sans-serif}
th,td{border:1px solid #cbd0d8;padding:5px 7px;vertical-align:top;text-align:left}th{background:#f3f4f6}
blockquote{margin:.8em 0;padding:.4em .9em;border-left:3px solid #cbd0d8;background:#f7f8fa}
.src{font-family:-apple-system,Helvetica,Arial,sans-serif;font-size:8.5pt;color:#4b5563;background:#f7f8fa;border-left:3px solid #e5e7eb;padding:4px 8px;margin:.2em 0 .6em}
.src b{color:#1f4e8c}
.meta{font-family:-apple-system,Helvetica,Arial,sans-serif;font-size:9pt;color:#6b7280;margin:0 0 1.4em}
hr{border:0;border-top:1px solid #e5e7eb;margin:1.4em 0}
@page{margin:.75in .7in}
"""


def strip_blocks_for_print(md):
    """Remove :::why blocks and ::: markers; turn :::details titles into headings."""
    out, lines, i = [], md.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if re.match(r"^:::why\s", ln):
            while i < len(lines) and lines[i].strip() != ":::":
                i += 1
            i += 1
            continue
        if ln.startswith(":::details "):
            out.append("### " + ln[len(":::details "):].strip()); i += 1; continue
        if ln.strip() in (":::register", ":::"):
            i += 1; continue
        out.append(ln); i += 1
    return "\n".join(out)


def convert_print(md, with_sources):
    out, list_stack, lines, i = [], [], md.splitlines(), 0
    heading_srcs = None

    def close_lists():
        while list_stack:
            out.append(f"</{list_stack.pop()}>")

    def flush_srcs():
        nonlocal heading_srcs
        if with_sources and heading_srcs:
            out.append('<div class="src">' + " ".join(f"<b>{html.escape(l)}:</b> {inline(t)}" for l, t in heading_srcs) + "</div>")
        heading_srcs = None

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("::src "):
            if heading_srcs is not None:
                label, _, text = ln[6:].partition(" :: ")
                heading_srcs.append((label.strip(), text.strip()))
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            flush_srcs(); close_lists()
            lvl = len(m.group(1))
            out.append(f"<h{lvl}>{inline(m.group(2).strip())}</h{lvl}>")
            heading_srcs = []
            i += 1
            continue
        flush_srcs()
        if not ln.strip():
            close_lists(); i += 1; continue
        if ln.strip() == "---":
            close_lists(); out.append("<hr>"); i += 1; continue
        if ln.startswith("|"):
            close_lists()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            out.append(table_html(parse_table(rows)))
            continue
        if ln.startswith("> "):
            close_lists()
            q = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i][1:].strip()); i += 1
            paras = "\n".join(q).split("\n\n")
            out.append("<blockquote>" + "".join(f"<p>{inline(p.replace(chr(10), ' '))}</p>" for p in paras if p.strip()) + "</blockquote>")
            continue
        mb = re.match(r"^(\s*)- (.*)$", ln)
        mo = re.match(r"^(\s*)\d+\. (.*)$", ln)
        if mb or mo:
            kind = "ul" if mb else "ol"
            if not list_stack or list_stack[-1] != kind:
                close_lists(); out.append(f"<{kind}>"); list_stack.append(kind)
            out.append(f"<li>{inline((mb or mo).group(2))}</li>")
            i += 1
            continue
        p = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            p.append(lines[i].strip()); i += 1
        close_lists()
        out.append(f"<p>{inline(' '.join(p))}</p>")
    flush_srcs(); close_lists()
    return "\n".join(out)


def find_chrome():
    for name in ["Google Chrome", "Chromium", "Microsoft Edge"]:
        if pathlib.Path(f"/Applications/{name}.app").exists():
            return name
    return None


def make_pdf(html_text, out_pdf, chrome_app, workdir):
    """Print HTML to PDF with headless Chrome, launched through launchd (`open`) so it is not
    constrained by the calling shell's sandbox; waits for the file to appear."""
    import time
    src = workdir / (out_pdf.stem + ".html")
    src.write_text(html_text)
    if out_pdf.exists():
        out_pdf.unlink()
    profile = workdir / f"profile-{out_pdf.stem}"     # fresh profile per PDF: Chrome locks a profile per instance
    cmd = ["open", "-n", "-W", "-a", chrome_app, "--args", "--headless=new", "--disable-gpu", "--no-first-run",
           "--no-default-browser-check", "--use-mock-keychain", "--password-store=basic", "--disable-extensions",
           f"--user-data-dir={profile}", "--no-pdf-header-footer",
           f"--print-to-pdf={out_pdf}", src.as_uri()]
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    # Chrome writes the PDF then lingers; stop waiting once the file size has been stable for a second.
    deadline, last, stable = time.time() + 60, -1, 0
    while time.time() < deadline:
        size = out_pdf.stat().st_size if out_pdf.exists() else 0
        if size > 0 and size == last:
            stable += 1
            if stable >= 2:
                break
        else:
            stable = 0
        last = size
        time.sleep(0.5)
    try:
        proc.wait(timeout=3)
    except subprocess.TimeoutExpired:
        proc.kill()
    subprocess.run(["pkill", "-9", "-f", str(profile)], capture_output=True)
    if not (out_pdf.exists() and out_pdf.stat().st_size > 0):
        raise RuntimeError("no PDF produced within 60s")


def md_to_docx(md, out_path, title, meta):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(11)
    style_names = {s.name for s in doc.styles}

    def add_runs(par, text):
        for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", text):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                par.add_run(tok[2:-2]).bold = True
            elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
                par.add_run(tok[1:-1]).font.name = "Menlo"
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                par.add_run(tok[1:-1]).italic = True
            else:
                par.add_run(tok)

    doc.add_heading(title, 0)
    p = doc.add_paragraph(); r = p.add_run(meta); r.italic = True
    lines = strip_blocks_for_print(md).splitlines()
    i, seen_h1 = 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("::src ") or ln.strip() == "---":
            i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            if lvl == 1 and not seen_h1:
                seen_h1 = True
            else:
                doc.add_heading(m.group(2).strip(), max(1, min(lvl, 3)))
            i += 1; continue
        if not ln.strip():
            i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            cells = parse_table(rows)
            if cells:
                ncol = max(len(r) for r in cells)
                t = doc.add_table(rows=0, cols=ncol)
                if "Table Grid" in style_names:
                    t.style = "Table Grid"
                for ri, r in enumerate(cells):
                    row = t.add_row().cells
                    for ci in range(ncol):
                        row[ci].text = ""
                        par = row[ci].paragraphs[0]
                        add_runs(par, r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in par.runs:
                                run.bold = True
                doc.add_paragraph("")
            continue
        if ln.startswith("> "):
            q = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i][1:].strip()); i += 1
            for para in "\n".join(q).split("\n\n"):
                if para.strip():
                    par = doc.add_paragraph(style="Quote") if "Quote" in style_names else doc.add_paragraph()
                    add_runs(par, para.replace("\n", " "))
            continue
        mb = re.match(r"^(\s*)- (.*)$", ln)
        mo = re.match(r"^(\s*)\d+\. (.*)$", ln)
        if mb or mo:
            sty = "List Bullet" if mb else "List Number"
            par = doc.add_paragraph(style=sty) if sty in style_names else doc.add_paragraph()
            add_runs(par, (mb or mo).group(2)); i += 1; continue
        pl = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            pl.append(lines[i].strip()); i += 1
        add_runs(doc.add_paragraph(), " ".join(pl))
    doc.save(out_path)


def build_downloads(ver, date, chrome):
    dl = SITE / "downloads"
    dl.mkdir(exist_ok=True)
    work = pathlib.Path(tempfile.mkdtemp(prefix="chss-pdf-"))
    results = {}
    for key, title, fn, group, base in DOCS:
        if not base:
            continue
        md = (ROOT / "content" / fn).read_text()
        h1 = re.search(r"^#\s+(.*)$", md, re.M)
        doc_title = h1.group(1).strip() if h1 else title
        meta = f"Draft v{ver} · {date} · Center for High School Success · for discussion; not legal advice"
        files = {}
        printable = strip_blocks_for_print(md)
        for variant, with_src in (("", False), ("-with-sources", True)):
            body = convert_print(printable, with_src)
            page = (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(doc_title)}</title>"
                    f"<style>{PRINT_CSS}</style></head><body><p class='meta'>{html.escape(meta)}</p>{body}</body></html>")
            if chrome:
                out_pdf = dl / f"{base}{variant}.pdf"
                try:
                    make_pdf(page, out_pdf, chrome, work)
                    files["pdf" + variant] = out_pdf.name
                except Exception as e:
                    print("pdf skipped:", base, variant, e, file=sys.stderr)
        try:
            md_to_docx(md, dl / f"{base}.docx", doc_title, meta)
            files["docx"] = f"{base}.docx"
        except Exception as e:
            print("docx skipped:", base, e, file=sys.stderr)
        (dl / f"{base}.md").write_text(md)
        files["md"] = f"{base}.md"
        results[key] = files
    shutil.rmtree(work, ignore_errors=True)
    return results


def download_row(files):
    if not files:
        return ""
    labels = [("pdf", "PDF"), ("pdf-with-sources", "PDF with sources"), ("docx", "Word"), ("md", "Markdown")]
    links = [f'<a class="dl" href="downloads/{files[k]}" download>{t}</a>' for k, t in labels if k in files]
    return '<div class="dlrow"><span class="dl-l">Download</span>' + "".join(links) + "</div>"


# ---------------------------------------------------------------- main
def version_info():
    raw = (ROOT / "VERSION").read_text().strip().splitlines()
    return raw[0].strip(), (raw[1].strip() if len(raw) > 1 else "")


def versions_list(current):
    vdir = SITE / "versions"
    found = set()
    if vdir.exists():
        for d in vdir.iterdir():
            if d.is_dir() and (d / "index.html").exists():
                found.add(d.name.lstrip("v"))
    found.add(current)
    return sorted(found, key=lambda v: tuple(int(x) if x.isdigit() else 0 for x in v.split(".")), reverse=True)


def main():
    tpl = (ROOT / "template.html").read_text()
    ver, date = version_info()
    chrome = find_chrome()
    downloads = build_downloads(ver, date, chrome)
    panels, tabs, notetabs = [], [], []
    for key, title, fn, group, base in DOCS:
        body = convert((ROOT / "content" / fn).read_text(), title_extra=download_row(downloads.get(key)))
        panels.append(f'<article class="doc" id="doc-{key}" data-doc="{key}" hidden>{body}</article>')
        t = f'<button class="tab" type="button" data-doc="{key}" role="tab" aria-selected="false">{html.escape(title)}</button>'
        (notetabs if group == "note" else tabs).append(t)

    vlist = versions_list(ver)
    vitems = ['<li><a href="#changelog">What changed (changelog)</a></li>']
    for v in vlist:
        label = f"v{v}" + (" (current)" if v == ver else "")
        href = "./" if v == ver else f"versions/v{v}/"
        vitems.append(f'<li><a href="{href}">{label}</a></li>')

    page = (tpl.replace("<!--TABS-->", "\n".join(tabs))
               .replace("<!--NOTETABS-->", "\n".join(notetabs))
               .replace("<!--PANELS-->", "\n".join(panels))
               .replace("<!--VERSIONS-->", "".join(vitems))
               .replace("{{VERSION}}", html.escape(ver))
               .replace("{{DATE}}", html.escape(date)))
    (SITE / "index.html").write_text(page)

    snap_dir = SITE / "versions" / f"v{ver}"
    snap_dir.mkdir(parents=True, exist_ok=True)
    snap = (page.replace('href="assets/', 'href="../../assets/')
                .replace('href="versions/', 'href="../../versions/')
                .replace('href="./"', 'href="../../"'))
    (snap_dir / "index.html").write_text(snap)
    if (SITE / "downloads").exists():
        shutil.copytree(SITE / "downloads", snap_dir / "downloads", dirs_exist_ok=True)
    names = sorted({n for f in downloads.values() for n in f.values()})
    print(f"wrote {SITE / 'index.html'} ({len(page)} bytes) and {snap_dir}; downloads: {names}; versions: {vlist}")


if __name__ == "__main__":
    main()
